"""Gera data/base_dados_biomateriais.xlsx a partir de data/*.csv e data/leia_me.md.

Os CSV são a fonte de verdade; o .xlsx é só para consulta: tem valores (sem fórmulas),
cabeçalhos formatados, painéis fixos e filtros. Não editar o .xlsx à mão.
Uso: python scripts/exportar_xlsx.py
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from biomat_mcdm.indices import SIGMA_E, stiffness_ratio
from biomat_mcdm.io import DATA, load_data, load_tissue

NOTA = "Gerado automaticamente a partir de data/*.csv; não editar à mão."
CABECALHO = Font(bold=True)
FUNDO = PatternFill("solid", fgColor="E8EEF6")


def _folha_tabela(wb: Workbook, nome: str, df: pd.DataFrame, filtros: bool = True) -> None:
    """Escreve uma tabela: cabeçalho a negrito com fundo, painel fixo, filtros, larguras."""
    ws = wb.create_sheet(nome)
    ws.append(list(df.columns))
    for linha in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in linha])
    for c in ws[1]:
        c.font, c.fill = CABECALHO, FUNDO
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if filtros and ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for j, col in enumerate(df.columns, start=1):
        largura = max([len(str(col))] + [len(str(v)) for v in df[col].dropna().head(300)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, largura + 2), 60)


def leia_me_linhas(data_dir: Path = DATA) -> list[str]:
    """Texto da folha LEIA-ME: a nota de geração automática e o conteúdo de leia_me.md."""
    texto = (data_dir / "leia_me.md").read_text(encoding="utf-8").splitlines()
    return [NOTA, ""] + [l.lstrip("#").strip() for l in texto]


def soma_pesos(crit: pd.DataFrame) -> pd.DataFrame:
    """Soma dos pesos por caso (verificação; deve ser 1, stent 0,85 no cenário A)."""
    s = crit.dropna(subset=["peso"]).groupby("caso_componente", sort=False)["peso"].sum()
    return s.round(4).reset_index().rename(columns={"peso": "soma_dos_pesos"})


def construir(data_dir: Path = DATA) -> Workbook:
    """Constrói o livro completo em memória (sem gravar)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "LEIA-ME"
    for l in leia_me_linhas(data_dir):
        ws.append([l])
    ws["A1"].font = Font(bold=True, color="B00020")
    ws.column_dimensions["A"].width = 120

    mat = pd.read_csv(data_dir / "materiais.csv")
    crit = pd.read_csv(data_dir / "criterios.csv")
    _folha_tabela(wb, "Materiais", mat)
    _folha_tabela(wb, "Criterios", crit)
    _folha_tabela(wb, "Soma_pesos", soma_pesos(crit), filtros=False)

    completo, _ = load_data(data_dir)
    sig = completo[completo["propriedade"] == SIGMA_E][["material", "min", "max", "tipico"]]
    sig = sig.rename(columns={"min": "sigma_y_E_min", "max": "sigma_y_E_max", "tipico": "sigma_y_E_tipico"})
    _folha_tabela(wb, "Indices_stent", sig.round(6), filtros=False)
    _folha_tabela(wb, "Indices_haste", stiffness_ratio(completo, load_tissue(data_dir)).round(3), filtros=False)

    _folha_tabela(wb, "Cenarios", pd.read_csv(data_dir / "cenarios.csv"), filtros=False)
    _folha_tabela(wb, "Tecido", pd.read_csv(data_dir / "tecido.csv"))
    _folha_tabela(wb, "Parametros", pd.read_csv(data_dir / "parametros.csv"), filtros=False)
    _folha_tabela(wb, "Problemas_criticos", pd.read_csv(data_dir / "problemas_criticos.csv"))
    return wb


def main() -> None:
    destino = DATA / "base_dados_biomateriais.xlsx"
    construir().save(destino)
    print(f"gravado: {destino}")


if __name__ == "__main__":
    main()
